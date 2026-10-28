"""Period inputs and analytic discounted cash-flow regression cases."""
from dataclasses import replace
from pathlib import Path
import tempfile
import unittest
import ValidadorFramework2 as v


class PeriodInputTests(unittest.TestCase):
    def test_scalar_and_exact_list(self):
        self.assertEqual(v.period_values('3', 3, 'test'), (3, 3, 3))
        self.assertEqual(v.period_values('0, 2,3', 3, 'test'), (0, 2, 3))

    def test_invalid_lists(self):
        for text in ('', '1,2', '1,2,3,4', '1,2,', '1,,3', '-1', 'nan', 'inf', '2x'):
            with self.subTest(text=text), self.assertRaises(ValueError):
                v.period_values(text, 3, 'test')
        with self.assertRaises(ValueError):
            v.period_values('0', 3, 'rate', positive=True)

    def test_all_fields_and_default_investment(self):
        source = Path(__file__).parent / 'examples/configuration.txt'
        self.assertEqual(v.load_configuration(source).investment, (0,))
        text = source.read_text().replace('NumberOfPeriods=1', 'NumberOfPeriods=3')
        text = text.replace('MiningCapacity=3', 'MiningCapacity=1,2,0')
        text = text.replace('MiningCost=1', 'MiningCost=1,3,5')
        text = text.replace('ProcessingCapacity=1.5', 'ProcessingCapacity=0.5,0.25,0')
        text = text.replace('ProcessingRate=1;', 'ProcessingRate=1,2,3;')
        text = text.replace('ProcessingCost=0;', 'ProcessingCost=2,4,6;')
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'config.txt'
            path.write_text(text + '\nInvestment=100,11,0\n')
            cfg = v.load_configuration(path)
        self.assertEqual(cfg.investment, (100, 11, 0))
        self.assertEqual(cfg.mining_capacity, (1, 2, 0))
        self.assertEqual(cfg.mining_cost, (1, 3, 5))
        self.assertEqual(cfg.processing_capacity, (.5, .25, 0))
        self.assertEqual(cfg.modes[0].rate, (1, 2, 3))
        self.assertEqual(cfg.modes[0].cost, (2, 4, 6))


class PeriodSolverTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import gurobipy as gp
        cls.gp = gp
        cls.env = gp.Env(empty=True)
        cls.env.setParam('OutputFlag', 0)
        cls.env.start()

    @classmethod
    def tearDownClass(cls):
        cls.env.dispose()

    def test_variable_parameters_and_empty_investment(self):
        base = v.load_configuration(Path(__file__).parent / 'examples/configuration.txt')
        cfg = replace(base, periods=3, discount=.1, rocks=('R1',), prices=(20,),
                      investment=(100, 11, 0), mining_cost=(1, 3, 5), mining_capacity=(1, 2, 0),
                      processing_capacity=(.5, .25, 0),
                      modes=(v.Mode('A', (1, 2, 3), (2, 4, 6), (1,), (1,)),
                             v.Mode('B', (1, 1, 1), (10, 10, 10), (1,), (1,))))
        block = v.Block(1, (1,), 'R1')
        expected = (-92, (8 - 3 - 11) / 1.1, -5 / 1.21)
        rows = []
        for t in range(3):
            row = v.solve_period(cfg, t, [block], self.env, self.gp)
            rows.append(row)
            self.assertAlmostEqual(row['PeriodTotal'], expected[t])
            self.assertAlmostEqual(row['ProcessingHours'], cfg.processing_capacity[t])
            self.assertAlmostEqual(row['DiscountedProcessingCost'], (1, 2 / 1.1, 0)[t])
            empty = v.solve_period(cfg, t, [], self.env, self.gp)
            self.assertAlmostEqual(empty['PeriodTotal'], -cfg.investment[t] / 1.1 ** t)
            without = v.solve_period(replace(cfg, investment=(0, 0, 0)), t, [block], self.env, self.gp)
            self.assertAlmostEqual(without['TotalOreProcessed'], row['TotalOreProcessed'])
            self.assertAlmostEqual(without['PeriodTotal'] - row['PeriodTotal'], row['DiscountedInvestment'])
        # Unit prices/rates must not be summed in Excel; cash flows must reconcile.
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'periods.xlsx'
            v.write_report(path, cfg, {0: rows}, {0: []}, {})
            wb = load_workbook(path)
            self.assertAlmostEqual(wb['Resumen']['D2'].value, sum(expected))
            detail = wb['Resultados']
            headers = [cell.value for cell in detail[1]]
            self.assertIsNone(detail.cell(5, headers.index('ProcessingRate_ModeA') + 1).value)
            self.assertIsNone(detail.cell(5, headers.index('MiningCost') + 1).value)
            self.assertAlmostEqual(detail.cell(5, headers.index('DiscountedInvestment') + 1).value, 110)
            wb.close()


if __name__ == '__main__':
    unittest.main()
