"""Run from repository root: python -m unittest discover -s validation -v."""
import contextlib
import io
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

import ValidadorFramework2 as v

EXAMPLES = Path(__file__).parent / "examples"


class InputTests(unittest.TestCase):
    def setUp(self):
        self.cfg = v.load_configuration(EXAMPLES / "configuration.txt")

    def test_current_repository_configuration(self):
        cfg = v.load_configuration(v.ROOT / "configuration/configuration.txt")
        self.assertEqual(tuple(m.name for m in cfg.modes), ("1", "2"))
        self.assertEqual(cfg.prices, (13228, 0))

    def test_unordered_keys_and_comments(self):
        lines = (EXAMPLES / "configuration.txt").read_text().splitlines()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.txt"
            path.write_text("\n".join(reversed(lines)), encoding="utf-8")
            cfg = v.load_configuration(path)
        self.assertEqual(cfg.prices, (10, 20, 30))
        self.assertEqual(cfg.modes[0].name, "Third")

    def test_duplicate_plan_and_period_bounds(self):
        for body in ("0,1,1,1\n0,1,1,1\n", "1,1,1,1\n", "0,4,1,1\n"):
            with self.subTest(body=body), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "plan.txt"
                path.write_text("Period,XIndex,YIndex,ZIndex\n" + body)
                with self.assertRaises(ValueError):
                    v.load_plan(path, self.cfg)

    def test_optional_rock_and_malformed_scenario(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scenario.txt"
            path.write_text("1,1,1,1,2,3,4\n")
            self.assertEqual(v.load_scenario(path, self.cfg)[1, 1, 1].rock, "R1")
            path.write_text("1,1,1,1,2\n")
            with self.assertRaises(ValueError):
                v.load_scenario(path, self.cfg)

    def test_precedence_same_period_allowed(self):
        cfg = replace(self.cfg, bounds=((1, 1), (1, 1), (1, 2)), periods=2)
        self.assertTrue(v.precedence_issues({(1, 1, 1): 0}, cfg))
        self.assertTrue(v.precedence_issues({(1, 1, 1): 0, (1, 1, 2): 1}, cfg))
        self.assertFalse(v.precedence_issues({(1, 1, 1): 0, (1, 1, 2): 0}, cfg))

    def test_missing_scenario_fails_before_solver(self):
        with contextlib.redirect_stderr(io.StringIO()) as errors:
            code = v.main([str(EXAMPLES / "configuration.txt"), "0", "2", "--plan", str(EXAMPLES / "plan.txt")])
        self.assertEqual(code, 1)
        self.assertIn("Faltan escenarios", errors.getvalue())

    def test_percentiles_small_range_and_ties(self):
        results = {i: [{"PeriodTotal": i}] for i in range(20)}
        self.assertEqual(v.percentiles(results), {"P10": 1, "P50": 9, "P90": 17})
        self.assertEqual(v.percentiles({5: [{"PeriodTotal": 0}]}), {"P10": 5, "P50": 5, "P90": 5})


class SolverTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import gurobipy as gp
        cls.gp = gp
        cls.env = gp.Env(empty=True)
        cls.env.setParam("OutputFlag", 0)
        cls.env.start()

    @classmethod
    def tearDownClass(cls):
        cls.env.dispose()

    def setUp(self):
        self.cfg = v.load_configuration(EXAMPLES / "configuration.txt")

    def solve(self, blocks, cfg=None, period=0):
        return v.solve_period(cfg or self.cfg, period, blocks, self.env, self.gp)

    def test_three_metals_rocks_modes_fractional_processing(self):
        blocks = list(v.load_scenario(EXAMPLES / "scenario00.txt", self.cfg).values())
        row = self.solve(blocks)
        self.assertAlmostEqual(row["PeriodTotal"], 147)
        self.assertAlmostEqual(row["ProcessedMetal2"], 1.5)
        self.assertAlmostEqual(row["ProcessedMetal3"], 4)
        self.assertAlmostEqual(row["ProcessingHours"], 1.5)
        self.assertAlmostEqual(row["Ore_ModeFirst"], 0)

    def test_three_rock_blending_ratios_and_inactive_mode(self):
        mixed = v.Mode("Mix", (1,), (0,), (1, 1, 1), (2, 1, 1))
        cfg = replace(self.cfg, modes=(mixed,), processing_capacity=(4,))
        blocks = [v.Block(1, (10, 0, 0), r) for r in cfg.rocks]
        row = self.solve(blocks, cfg)
        self.assertAlmostEqual(row["R1_ModeMix"], 1)
        self.assertAlmostEqual(row["R2_ModeMix"], .5)
        self.assertAlmostEqual(row["R3_ModeMix"], .5)
        self.assertAlmostEqual(row["ProcessedMetal1"], 20)
        absent = self.solve(blocks[:2], cfg)
        self.assertAlmostEqual(absent["TotalOreProcessed"], 0)
        self.assertAlmostEqual(absent["PeriodTotal"], -2)

    def test_block_cannot_be_processed_twice(self):
        cfg = replace(self.cfg, processing_capacity=(10,),
                      modes=(self.cfg.modes[0], replace(self.cfg.modes[0], name="Other")))
        row = self.solve([v.Block(1, (2, 0, 0), "R1")], cfg)
        self.assertAlmostEqual(row["TotalOreProcessed"], 1)
        self.assertAlmostEqual(row["ProcessedMetal1"], 2)

    def test_actual_mass_cost_uniform_processing_mass_and_discount(self):
        cfg = replace(self.cfg, periods=2, investment=(0, 0), mining_capacity=(3, 3), mining_cost=(1, 1),
                      modes=(replace(self.cfg.modes[0], cost=(2, 2), rate=(1, 1)),), processing_capacity=(.5, .5))
        row = self.solve([v.Block(4, (2, 0, 0), "R1")], cfg, period=1)
        self.assertAlmostEqual(row["DiscountedValueNoMiningCost"], 6 / 1.08)
        self.assertAlmostEqual(row["DiscountedMiningCost"], 4 / 1.08)
        self.assertAlmostEqual(row["TotalOreProcessed"], .5)

    def test_empty_and_negative_value(self):
        self.assertEqual(self.solve([])["PeriodTotal"], 0)
        row = self.solve([v.Block(1, (0, 0, 0), "R1")], replace(
            self.cfg, modes=(replace(self.cfg.modes[0], cost=(2,)),)))
        self.assertAlmostEqual(row["TotalOreProcessed"], 0)
        self.assertAlmostEqual(row["PeriodTotal"], -1)

    def test_missing_block_uses_defaults(self):
        cfg = replace(self.cfg, default_metal=1)
        rows = v.evaluate(cfg, {(1, 1, 1): 0}, {}, self.env, self.gp)
        self.assertAlmostEqual(rows[0]["PeriodTotal"], 59)

    def test_cli_report_and_infeasible_capacity(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            out = Path(directory) / "report.xlsx"
            with contextlib.redirect_stdout(io.StringIO()):
                code = v.main([str(EXAMPLES / "configuration.txt"), "0", "1", "--plan",
                               str(EXAMPLES / "plan.txt"), "--out", str(out), "--quiet"])
            self.assertEqual(code, 0)
            wb = load_workbook(out)
            self.assertEqual(wb["Resumen"]["D2"].value, 147)
            self.assertEqual(wb["Resumen"]["D3"].value, 87)
            self.assertEqual(len(wb["Cumulative NPV"]._charts), 1)
            self.assertEqual(wb["Comprobaciones"].max_row, 1)
            wb.close()
            config = Path(directory) / "configuration.txt"
            config.write_text((EXAMPLES / "configuration.txt").read_text().replace("MiningCapacity=3", "MiningCapacity=2"))
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                code = v.main([str(config), "0", "--plan", str(EXAMPLES / "plan.txt"), "--out", str(out), "--quiet"])
            self.assertEqual(code, 2)
            wb = load_workbook(out)
            self.assertFalse(wb["Resumen"]["C2"].value)
            self.assertIn("Capacidad minera", wb["Comprobaciones"]["C2"].value)
            wb.close()


if __name__ == "__main__":
    unittest.main()
