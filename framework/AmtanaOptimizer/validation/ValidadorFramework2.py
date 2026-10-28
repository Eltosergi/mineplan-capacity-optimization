#!/usr/bin/env python3
"""Evaluate a fixed AmtanaOptimizer plan with an exact processing LP per scenario.

Run --help for arguments; see README.md for model conventions and setup.
Gurobi and Excel dependencies are loaded only when an evaluation is requested.
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def number(value):
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"Valor no finito: {value}")
    return result


def vector(value):
    return tuple(number(x) for x in value.split(","))


def period_values(value, periods, name, positive=False):
    result = vector(value)
    if periods <= 0 or len(result) not in (1, periods):
        raise ValueError(f"{name}: se requiere un valor o exactamente {periods} valores")
    if any(x < 0 or (positive and x == 0) for x in result):
        raise ValueError(f"{name}: valores deben ser {'positivos' if positive else 'no negativos'}")
    return result * periods if len(result) == 1 else result


@dataclass(frozen=True)
class Mode:
    name: str
    rate: tuple[float, ...]
    cost: tuple[float, ...]
    recoveries: tuple[float, ...]
    weights: tuple[float, ...]


@dataclass
class Config:
    values: dict[str, str]
    rocks: tuple[str, ...]
    prices: tuple[float, ...]
    modes: tuple[Mode, ...]
    periods: int
    discount: float
    default_tonnage: float
    default_metal: float
    default_rock: str
    mining_cost: tuple[float, ...]
    mining_capacity: tuple[float, ...]
    processing_capacity: tuple[float, ...]
    investment: tuple[float, ...]
    bounds: tuple[tuple[int, int], ...]


def load_configuration(path):
    values, raw_modes = {}, []
    for line_no, line in enumerate(Path(path).read_text(encoding="utf-8-sig").splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        key, sep, value = line.partition("=")
        if not sep:
            raise ValueError(f"{path}:{line_no}: se esperaba Clave=Valor")
        key, value = key.strip(), value.strip()
        if key == "OperationalMode":
            raw_modes.append(value)
        else:
            if key in values:
                raise ValueError(f"Clave duplicada: {key}")
            values[key] = value
    required = ("FileNameBase FileNameExtension MinX MaxX MinY MaxY MinZ MaxZ "
                "DefaultRockTonnage DefaultMetalTonnage DefaultRockType NumberOfPeriods "
                "AnnualDiscountRate MiningCapacity MiningCost NumberOfMetals "
                "EffectiveMetalPrices ProcessingCapacity NumberOfRockTypes RockTypes "
                "NumberOfOperationalModes").split()
    missing = set(required) - values.keys()
    if missing:
        raise ValueError("Faltan claves del formato actual: " + ", ".join(sorted(missing)))
    rocks = tuple(x.strip() for x in values["RockTypes"].split(","))
    prices = vector(values["EffectiveMetalPrices"])
    if (len(rocks) != int(values["NumberOfRockTypes"]) or not all(rocks)
            or len(set(rocks)) != len(rocks)):
        raise ValueError("RockTypes debe contener nombres únicos y coincidir con NumberOfRockTypes")
    if len(prices) != int(values["NumberOfMetals"]) or not prices:
        raise ValueError("EffectiveMetalPrices no coincide con NumberOfMetals")
    periods = int(values["NumberOfPeriods"])
    modes = []
    for raw in raw_modes:
        fields = {}
        for item in raw.split(";"):
            key, sep, value = item.partition("=")
            if not sep or key.strip() in fields:
                raise ValueError(f"Campo de modo inválido o duplicado: {item}")
            fields[key.strip()] = value.strip()
        if not {"Name", "ProcessingRate", "ProcessingCost", "Recoveries", "w"} <= fields.keys():
            raise ValueError("Modo incompleto: se requiere Name, ProcessingRate, ProcessingCost, Recoveries, w")
        mode = Mode(fields["Name"], period_values(fields["ProcessingRate"], periods, "ProcessingRate", True),
                    period_values(fields["ProcessingCost"], periods, "ProcessingCost"), vector(fields["Recoveries"]), vector(fields["w"]))
        if (not mode.name or len(mode.recoveries) != len(prices)
                or len(mode.weights) != len(rocks) or min(mode.weights) < 0
                or max(mode.weights) <= 0):
            raise ValueError(f"Parámetros inválidos para el modo {mode.name}")
        modes.append(mode)
    if not modes or len(modes) != int(values["NumberOfOperationalModes"]):
        raise ValueError("El número de líneas OperationalMode no coincide con NumberOfOperationalModes")
    if len({m.name for m in modes}) != len(modes):
        raise ValueError("Los nombres de modos deben ser únicos")
    cfg = Config(values, rocks, prices, tuple(modes), int(values["NumberOfPeriods"]),
                 number(values["AnnualDiscountRate"]), number(values["DefaultRockTonnage"]),
                 number(values["DefaultMetalTonnage"]), values["DefaultRockType"],
                 period_values(values["MiningCost"], periods, "MiningCost"),
                 period_values(values["MiningCapacity"], periods, "MiningCapacity"),
                 period_values(values["ProcessingCapacity"], periods, "ProcessingCapacity"),
                 period_values(values.get("Investment", "0"), periods, "Investment"),
                 tuple((int(values[f"Min{a}"]), int(values[f"Max{a}"])) for a in "XYZ"))
    if (cfg.periods <= 0 or cfg.discount <= -1 or cfg.default_tonnage <= 0
            or cfg.default_rock not in rocks or any(lo > hi for lo, hi in cfg.bounds)):
        raise ValueError("Horizonte, límites, capacidades o valores por defecto inválidos")
    return cfg


def in_bounds(coord, cfg):
    return all(lo <= x <= hi for x, (lo, hi) in zip(coord, cfg.bounds))


def load_plan(path, cfg):
    plan = {}
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        if next(reader, None) != ["Period", "XIndex", "YIndex", "ZIndex"]:
            raise ValueError("Encabezado del plan: Period,XIndex,YIndex,ZIndex")
        for row in reader:
            if not row:
                continue
            if len(row) != 4:
                raise ValueError(f"Plan, línea {reader.line_num}: se esperan 4 columnas")
            period, x, y, z = map(int, row)
            coord = (x, y, z)
            if coord in plan or not in_bounds(coord, cfg) or not 0 <= period < cfg.periods:
                raise ValueError(f"Plan, línea {reader.line_num}: bloque duplicado, fuera del modelo o período inválido")
            plan[coord] = period
    return plan


def precedence_issues(plan, cfg):
    issues = []
    for (x, y, z), period in plan.items():
        if z == cfg.bounds[2][1]:
            continue
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                above = (x + dx, y + dy, z + 1)
                if in_bounds(above, cfg) and plan.get(above, cfg.periods) > period:
                    issues.append((period, f"Bloque {(x, y, z)} requiere {above} en el mismo período o antes"))
    return issues


@dataclass(frozen=True)
class Block:
    tonnage: float
    metals: tuple[float, ...]
    rock: str


def load_scenario(path, cfg):
    blocks = {}
    k = len(cfg.prices)
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        for row in reader:
            if not row:
                continue
            if len(row) not in (4 + k, 5 + k):
                raise ValueError(f"{path}:{reader.line_num}: se esperan X,Y,Z,RockTonnage,{k} metales[,RockType]")
            coord = tuple(map(int, row[:3]))
            block = Block(number(row[3]), tuple(number(x) for x in row[4:4+k]),
                          row[4+k].strip() if len(row) == 5+k else cfg.default_rock)
            if (coord in blocks or not in_bounds(coord, cfg)
                    or block.rock not in cfg.rocks or block.tonnage < 0):
                raise ValueError(f"{path}:{reader.line_num}: coordenadas, tonelaje o roca inválidos")
            blocks[coord] = block
    return blocks


def solve_period(cfg, period, blocks, env, gp):
    """Continuous block fractions; blending uses uniform default tonnage, as in C++."""
    discount = (1 + cfg.discount) ** period
    row = {"Period": period, "DiscountedValueNoMiningCost": 0.0,
           "DiscountedMiningCost": sum(b.tonnage for b in blocks) * cfg.mining_cost[period] / discount,
           "RockExtracted": sum(b.tonnage for b in blocks), "TotalOreProcessed": 0.0,
           "ProcessingHours": 0.0,
           "Investment": cfg.investment[period],
           "DiscountedInvestment": cfg.investment[period] / discount,
           "MiningCapacity": cfg.mining_capacity[period],
           "MiningCost": cfg.mining_cost[period],
           "ProcessingCapacity": cfg.processing_capacity[period],
           "DiscountedProcessingCost": 0.0}
    for rock in cfg.rocks:
        row[f"RockExtracted_{rock}"] = sum(b.tonnage for b in blocks if b.rock == rock)
    for j in range(len(cfg.prices)):
        row[f"ProcessedMetal{j+1}"] = 0.0
    for mode in cfg.modes:
        row[f"Ore_Mode{mode.name}"] = 0.0
        row[f"ProcessingRate_Mode{mode.name}"] = mode.rate[period]
        row[f"ProcessingCost_Mode{mode.name}"] = mode.cost[period]
        for rock in cfg.rocks:
            row[f"{rock}_Mode{mode.name}"] = 0.0
    if blocks:
        with gp.Model(env=env) as model:
            model.Params.Method = 0
            present = {b.rock for b in blocks}
            keys = [(i, m) for m, mode in enumerate(cfg.modes)
                    if all(rock in present for rock, w in zip(cfg.rocks, mode.weights) if w > 0)
                    for i, block in enumerate(blocks) if mode.weights[cfg.rocks.index(block.rock)] > 0]
            fractions = model.addVars(keys, lb=0, ub=1, name="fraction")
            for i in range(len(blocks)):
                model.addConstr(gp.quicksum(fractions[i, m] for m in range(len(cfg.modes))
                                           if (i, m) in fractions) <= 1)
            model.addConstr(gp.quicksum(fractions[i, m] * cfg.default_tonnage / cfg.modes[m].rate[period]
                                       for i, m in keys) <= cfg.processing_capacity[period])
            for m, mode in enumerate(cfg.modes):
                active = [r for r, w in enumerate(mode.weights) if w > 0]
                ref = max(active, key=lambda r: mode.weights[r])
                amounts = {r: gp.quicksum(fractions[i, m] for i, b in enumerate(blocks)
                                          if b.rock == cfg.rocks[r] and (i, m) in fractions)
                           for r in active}
                for r in active:
                    if r != ref:
                        model.addConstr(amounts[r] == mode.weights[r] / mode.weights[ref] * amounts[ref])
            model.setObjective(gp.quicksum(
                fractions[i, m] * (sum(metal * recovery * price for metal, recovery, price in
                                      zip(blocks[i].metals, cfg.modes[m].recoveries, cfg.prices))
                                   - blocks[i].tonnage * cfg.modes[m].cost[period]) / discount
                for i, m in keys), gp.GRB.MAXIMIZE)
            model.optimize()
            if model.Status != gp.GRB.OPTIMAL:
                raise RuntimeError(f"Período {period}: Gurobi terminó con estado {model.Status}")
            row["DiscountedValueNoMiningCost"] = model.ObjVal
            for i, m in keys:
                fraction = fractions[i, m].X
                ore = fraction * cfg.default_tonnage
                row["TotalOreProcessed"] += ore
                row["DiscountedProcessingCost"] += fraction * blocks[i].tonnage * cfg.modes[m].cost[period] / discount
                row["ProcessingHours"] += ore / cfg.modes[m].rate[period]
                row[f"Ore_Mode{cfg.modes[m].name}"] += ore
                row[f"{blocks[i].rock}_Mode{cfg.modes[m].name}"] += ore
                for j, metal in enumerate(blocks[i].metals):
                    row[f"ProcessedMetal{j+1}"] += fraction * metal
    row["PeriodTotal"] = row["DiscountedValueNoMiningCost"] - row["DiscountedMiningCost"] - row["DiscountedInvestment"]
    return row


def evaluate(cfg, plan, scenario, env, gp):
    default = Block(cfg.default_tonnage, (cfg.default_metal,) * len(cfg.prices), cfg.default_rock)
    periods = [[] for _ in range(cfg.periods)]
    for coord, period in plan.items():
        periods[period].append(scenario.get(coord, default))
    return [solve_period(cfg, t, blocks, env, gp) for t, blocks in enumerate(periods)]


def percentiles(results):
    ordered = sorted(results, key=lambda idx: (sum(r["PeriodTotal"] for r in results[idx]), idx))
    return {label: ordered[max(1, math.ceil(p * len(ordered))) - 1]
            for label, p in (("P10", .1), ("P50", .5), ("P90", .9))}


def write_report(path, cfg, results, issues, metadata):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    pct = percentiles(results)
    def additive(key):
        return key not in ("Period", "MiningCost") and not key.startswith(("ProcessingRate_Mode", "ProcessingCost_Mode"))
    metrics = [key for key in next(iter(results.values()))[0] if additive(key)]
    summary.append(["Scenario", "Percentil", "PlanFeasible", "NPV_Final"] + metrics)
    detail = wb.create_sheet("Resultados" if len(results) == 1 else "Detalle")
    columns = list(next(iter(results.values()))[0])
    detail.append(["Scenario"] + columns)
    for idx, rows in sorted(results.items()):
        total = sum(r["PeriodTotal"] for r in rows)
        summary.append([idx, "/".join(p for p, s in pct.items() if s == idx), not issues[idx], total]
                       + [sum(r[k] for r in rows) for k in metrics])
        for row in rows:
            detail.append([idx] + [row[k] for k in columns])
        detail.append([idx] + ["TOTAL" if k == "Period" else sum(r[k] for r in rows) if additive(k) else None for k in columns])
    cumulative = wb.create_sheet("Cumulative NPV")
    curves = [(f"Scenario {next(iter(results))}", next(iter(results)))] if len(results) == 1 else list(pct.items())
    cumulative.append(["Period"] + [f"{label} (sc {idx})" for label, idx in curves])
    sums = [0.0] * len(curves)
    cumulative.append([0] + sums)
    for period in range(cfg.periods):
        sums = [value + results[idx][period]["PeriodTotal"] for value, (_, idx) in zip(sums, curves)]
        cumulative.append([period + 1] + sums)
    chart = LineChart()
    chart.title = "Cumulative NPV (Gurobi LP)"
    chart.x_axis.title, chart.y_axis.title = "Period", "NPV"
    chart.add_data(Reference(cumulative, min_col=2, max_col=len(curves)+1,
                             min_row=1, max_row=cfg.periods+2), titles_from_data=True)
    chart.set_categories(Reference(cumulative, min_col=1, min_row=2, max_row=cfg.periods+2))
    for series, (label, _) in zip(chart.series, curves):
        series.graphicalProperties.line.solidFill = "808080" if label in ("P10", "P90") else "1F4E78"
        if label in ("P10", "P90"):
            series.graphicalProperties.line.prstDash = "dash"
    chart.width, chart.height = 22, 11
    cumulative.add_chart(chart, f"{get_column_letter(len(curves)+3)}2")
    checks = wb.create_sheet("Comprobaciones")
    checks.append(["Scenario", "Period", "Issue"])
    for idx, found in sorted(issues.items()):
        for period, message in found:
            checks.append([idx, period, message])
    info = wb.create_sheet("Metadatos")
    info.append(["Campo", "Valor"])
    for key, value in metadata.items():
        info.append([key, str(value)])
    info.append(["Método", "LP exacto de procesamiento para un plan fijo; no reproduce la matheurística C++"])
    info.append(["Metales", "Contenido metálico procesado ponderado por fracción, antes de recuperación"])
    info.append(["Percentiles", "Nearest-rank sobre NPV final; P10 bajo, P90 alto; empates por índice"])
    info.append(["Factibilidad", "Precedencias de 9 vecinos superiores y capacidad minera por escenario; ver Comprobaciones"])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for column in sheet.columns:
            width = min(70, max(14, max(len(str(c.value or "")) for c in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.000000"
                # Input filenames and mode/rock names are text, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.data_type = "s"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def resolve(root, path):
    path = Path(path)
    return path.resolve() if path.is_absolute() else (root / path).resolve()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("configuration", help="configuración actual del optimizador")
    parser.add_argument("scenario_start", type=int)
    parser.add_argument("scenario_end", type=int, nargs="?", help="fin inclusivo; por defecto un escenario")
    parser.add_argument("--project-root", type=Path, default=ROOT, help="base de todas las rutas relativas; por defecto raíz del repositorio")
    parser.add_argument("--plan", default="sensor/result.txt")
    parser.add_argument("--out", help="reporte .xlsx; por defecto validation/results/validation_NN-MM.xlsx")
    parser.add_argument("--quiet", action="store_true", help="silenciar Gurobi")
    args = parser.parse_args(argv)
    end = args.scenario_end if args.scenario_end is not None else args.scenario_start
    if args.scenario_start < 0 or end < args.scenario_start:
        parser.error("Rango inválido: 0 <= inicio <= fin")
    try:
        root = args.project_root.resolve()
        config_path, plan_path = resolve(root, args.configuration), resolve(root, args.plan)
        cfg = load_configuration(config_path)
        plan = load_plan(plan_path, cfg)
        common_issues = precedence_issues(plan, cfg)
        paths = {idx: resolve(root, f"{cfg.values['FileNameBase']}{idx:02d}{cfg.values['FileNameExtension']}")
                 for idx in range(args.scenario_start, end + 1)}
        missing = [str(path) for path in paths.values() if not path.is_file()]
        if missing:
            raise ValueError("Faltan escenarios solicitados; no se omiten para calcular percentiles: " + ", ".join(missing))
        out = resolve(root, args.out or f"validation/results/validation_{args.scenario_start:02d}-{end:02d}.xlsx")
        if out.suffix.lower() != ".xlsx":
            raise ValueError("La salida debe tener extensión .xlsx")
        if out in {config_path, plan_path, *paths.values()}:
            raise ValueError("La salida no puede sobrescribir una entrada")
        import gurobipy as gp
        import openpyxl
        results, issues = {}, {}
        with gp.Env(empty=True) as env:
            env.setParam("OutputFlag", 0 if args.quiet else 1)
            env.start()
            for idx, path in paths.items():
                scenario = load_scenario(path, cfg)
                results[idx] = evaluate(cfg, plan, scenario, env, gp)
                issues[idx] = list(common_issues)
                for row in results[idx]:
                    capacity = cfg.mining_capacity[row["Period"]]
                    if row["RockExtracted"] > capacity + 1e-6 * max(1, capacity):
                        issues[idx].append((row["Period"], f"Capacidad minera excedida: {row['RockExtracted']} > {capacity}"))
                npv = sum(row["PeriodTotal"] for row in results[idx])
                print(f"Scenario {idx:02d}: NPV LP = {npv:.6f}; incidencias = {len(issues[idx])}")
        write_report(out, cfg, results, issues, {
            "Configuration": config_path, "Plan": plan_path, "ScenarioBase": resolve(root, cfg.values["FileNameBase"]),
            "ScenarioRange": f"{args.scenario_start}..{end}", "Python": sys.version.split()[0],
            "Gurobi": ".".join(map(str, gp.gurobi.version())), "openpyxl": openpyxl.__version__})
        print(f"Excel: {out}")
        if any(issues.values()):
            print("ADVERTENCIA: el plan presenta incidencias; consulte Comprobaciones.", file=sys.stderr)
            return 2
        return 0
    except ImportError as exc:
        print(f"ERROR: {exc}. Instale validation/requirements.txt en su entorno Python.", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
